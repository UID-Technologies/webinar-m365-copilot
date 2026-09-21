"""Generate all reference files required for Lab 01, Lab 02, and Lab 03."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent / "Reference Files"
LAB02 = ROOT / "Program Portfolio Knowledge"
LAB03 = ROOT / "Sprint Planning Knowledge"

TODAY = date(2026, 9, 21)


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def save_docx(doc: Document, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def style_header_row(ws, row: int = 1) -> None:
    fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font


def write_sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_header_row(ws)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18


def create_lab01_brd() -> None:
    doc = Document()
    title = doc.add_heading("Finance Transformation Business Requirements Document", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    para(doc, "Document Version: 1.0  |  Status: Draft for Training  |  Classification: Fictional / Sanitized")
    doc.add_paragraph()

    heading(doc, "1. Executive Summary")
    para(
        doc,
        "Contoso Finance is executing a multi-year Finance Transformation Program. Initiative "
        "information is spread across spreadsheets and documents. Senior finance leadership "
        "requires a single decision-support view to review the transformation portfolio, "
        "priorities, readiness, dependencies, risks, and data quality before investment discussions.",
    )

    heading(doc, "2. Business Problem")
    bullet(doc, "Transformation initiative data is fragmented across multiple sources.")
    bullet(doc, "Leadership lacks a consolidated portfolio view for decision discussions.")
    bullet(doc, "Important information is missing or inconsistent in current records.")
    bullet(doc, "Manual consolidation is slow and error-prone.")

    heading(doc, "3. Business Objectives")
    bullet(doc, "Provide an executive overview of the finance transformation portfolio.")
    bullet(doc, "Enable search, filter, and sort of initiatives using source-backed fields.")
    bullet(doc, "Support prioritization and roadmap review using explicit source data.")
    bullet(doc, "Surface dependencies, risks, and data-quality gaps.")
    bullet(doc, "Support leadership discussion—not automated investment approval.")

    heading(doc, "4. Primary Users")
    for user in [
        "CFO",
        "Finance Transformation Lead",
        "FP&A Lead",
        "Controller",
        "GBS Lead",
        "Process Owners",
        "Program Leadership",
    ]:
        bullet(doc, user)

    heading(doc, "5. Information Required")
    for item in [
        "KPI baselines",
        "Pain points",
        "Transformation initiatives",
        "Estimated costs",
        "Benefit ranges",
        "Readiness",
        "Dependencies",
        "Risks",
        "Target dates",
        "Initiative owners",
    ]:
        bullet(doc, item)

    heading(doc, "6. Functional Requirements")
    requirements = [
        ("FR-01", "Inspect the supplied dataset and map available columns to BRD requirements before building any solution."),
        ("FR-02", "Provide an Executive Overview with KPI cards supported only by source data."),
        ("FR-03", "Provide a searchable Initiative Portfolio."),
        ("FR-04", "Provide filters for source-backed categories such as function, status, priority, and readiness."),
        ("FR-05", "Provide a sortable initiative table."),
        ("FR-06", "Provide an initiative detail view."),
        ("FR-07", "Provide a prioritization view using only explicit source priority fields."),
        ("FR-08", "Provide a Transformation Roadmap using explicit target-date or timing information."),
        ("FR-09", "Provide a Dependencies and Risks view where supported by source data."),
        ("FR-10", "Provide a Data Quality section identifying missing or incomplete information."),
        ("FR-11", "Display a Professional Review notice stating the prototype is not an approved business case."),
        ("FR-12", "Reconcile KPI values to the supplied dataset."),
    ]
    for req_id, text in requirements:
        p = doc.add_paragraph(style="List Number")
        p.add_run(f"{req_id}: ").bold = True
        p.add_run(text)

    heading(doc, "7. Constraints and Business Rules")
    para(doc, "DO:", bold=True)
    bullet(doc, "Use only supplied source information.")
    bullet(doc, "Clearly identify missing information.")
    bullet(doc, "Keep source facts distinguishable from AI-generated observations.")
    bullet(doc, "Use fictional or approved sanitized data only.")
    para(doc, "DO NOT:", bold=True)
    bullet(doc, "Invent missing costs, benefits, owners, dates, or dependencies.")
    bullet(doc, "Create unsupported numerical scores or rankings.")
    bullet(doc, "Present the prototype as an approved business case or forecast.")
    bullet(doc, "Make investment recommendations.")
    para(doc, 'Unsupported information must be shown as "Not provided" or "Not available from supplied data."')

    heading(doc, "8. Expected Output")
    for item in [
        "Executive Overview",
        "KPI Cards",
        "Initiative Portfolio with Search / Filter / Sort",
        "Prioritization view",
        "Transformation Roadmap",
        "Dependencies & Risks",
        "Data Quality section",
        "Professional Review notice",
    ]:
        bullet(doc, item)

    heading(doc, "9. Acceptance Criteria")
    bullets = [
        "Information displayed is traceable to the source dataset or marked unavailable.",
        "Search returns only matching initiatives.",
        "Filters and sorting operate on source-backed fields only.",
        "Drill-down to initiative detail is supported.",
        "KPI cards reconcile to supplied data.",
        "Missing information is identified in the Data Quality section.",
        "Professional review notice is visible.",
    ]
    for item in bullets:
        bullet(doc, item)

    heading(doc, "10. Out of Scope")
    bullet(doc, "Approved business case or investment decision system.")
    bullet(doc, "Automated approval workflows.")
    bullet(doc, "Integration with live ERP or planning systems.")

    save_docx(doc, ROOT / "Finance_Transformation_BRD.docx")


def create_lab01_data() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Initiatives"

    headers = [
        "Initiative ID",
        "Initiative Name",
        "Function",
        "Status",
        "Priority",
        "Readiness",
        "Estimated Cost (USD)",
        "Benefit Range (USD)",
        "Owner",
        "Dependencies",
        "Risks",
        "Target Date",
        "KPI Baseline",
        "Pain Point",
    ]
    rows = [
        ["FT-001", "Close Process Automation", "Record to Report", "In Progress", "High", "Medium", 850000, "1.2M - 1.8M", "Sarah Chen", "FT-004", "Vendor delivery delay", "2026-03-31", "Close cycle 8 days", "Manual close steps"],
        ["FT-002", "FP&A Planning Platform", "FP&A", "Planning", "High", "Low", 1200000, "2.0M - 3.5M", "Michael Torres", "FT-006", "Data quality in source systems", "2026-06-30", "Planning cycle 15 days", "Spreadsheet-driven planning"],
        ["FT-003", "Intercompany Hub", "Controllership", "In Progress", "High", "Medium", 620000, "", "Lisa Park", "FT-001", "Cross-border regulatory change", "2026-04-15", "IC mismatch rate 2.1%", "Manual IC reconciliation"],
        ["FT-004", "Accounts Payable Workflow", "GBS", "In Progress", "High", "High", 430000, "0.8M - 1.1M", "David Okonkwo", "", "Change resistance", "2026-02-28", "Touchless rate 42%", "Paper-based approvals"],
        ["FT-005", "Management Reporting Hub", "FP&A", "Not Started", "Medium", "Low", 390000, "0.5M - 0.9M", "", "FT-002", "Unclear report ownership", "2026-09-30", "", "Fragmented reporting packs"],
        ["FT-006", "Data Governance Framework", "Finance IT", "In Progress", "Medium", "Medium", 275000, "Not quantified", "Priya Nair", "FT-002", "Limited data steward capacity", "2026-05-31", "Master data accuracy 88%", "Inconsistent chart of accounts"],
        ["FT-007", "Treasury Workbench", "Treasury", "Planning", "Low", "Low", 510000, "0.4M - 0.7M", "James Walsh", "", "Bank API availability", "2026-11-30", "Cash forecast accuracy 78%", "Manual cash forecasting"],
        ["FT-008", "Expense Policy Analytics", "GBS", "Complete", "Low", "High", 180000, "0.2M - 0.3M", "Emma Brooks", "", "", "2025-12-15", "Policy compliance 91%", "Limited spend visibility"],
    ]
    write_sheet(ws, headers, rows)

    ws2 = wb.create_sheet("KPI Summary")
    kpi_headers = ["KPI", "Baseline Value", "Target Value", "Notes"]
    kpi_rows = [
        ["Finance close cycle (days)", "8", "5", "Source: initiative FT-001"],
        ["Planning cycle (days)", "15", "7", "Source: initiative FT-002"],
        ["Touchless AP rate (%)", "42", "75", "Source: initiative FT-004"],
        ["Master data accuracy (%)", "88", "95", "Source: initiative FT-006"],
    ]
    write_sheet(ws2, kpi_headers, kpi_rows)

    path = ROOT / "Finance_Transformation_Data.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def create_lab02_brd() -> None:
    doc = Document()
    title = doc.add_heading("Program & Portfolio Intelligence Agent — Business Requirements Document", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    para(doc, "Document Version: 1.0  |  Status: Draft for Training  |  Classification: Fictional / Sanitized")
    doc.add_paragraph()

    heading(doc, "1. Purpose")
    para(
        doc,
        "Define requirements for a Program & Portfolio Intelligence Agent that helps PMO and "
        "program leadership consolidate approved program information and produce governance outputs.",
    )

    heading(doc, "2. Primary Users")
    for user in ["PMO", "Program Director", "Steering Committee", "Workstream Leads", "Finance", "Risk / Benefits Owners"]:
        bullet(doc, user)

    heading(doc, "3. Approved Knowledge Sources")
    for src in [
        "Program overview documents",
        "Project status updates",
        "RAID logs",
        "Action registers",
        "Decision registers",
        "Financial records",
        "Governance / Steering Committee minutes",
    ]:
        bullet(doc, src)

    heading(doc, "4. Functional Requirements")
    reqs = [
        ("AR-01", "Summarize project and portfolio status using configured knowledge sources."),
        ("AR-02", "Consolidate RAID items, actions, decisions, and milestones."),
        ("AR-03", "Identify overdue actions only when a recorded due date supports that conclusion."),
        ("AR-04", "Report dependencies only when explicitly recorded in approved sources."),
        ("AR-05", "Summarize recorded financial information preserving currency and reporting period."),
        ("AR-06", "Detect and report missing or conflicting information."),
        ("AR-07", "Generate governance summaries, status reports, RAID views, steering updates, decision briefs, and action registers."),
        ("AR-08", "Draft proposed actions and reminders without presenting them as approved or sent."),
        ("AR-09", "Use only configured sources and information the current user is authorized to access."),
        ("AR-10", "Make material statements traceable to configured sources."),
    ]
    for req_id, text in reqs:
        p = doc.add_paragraph(style="List Number")
        p.add_run(f"{req_id}: ").bold = True
        p.add_run(text)

    heading(doc, "5. Controls — Agent Must NOT")
    for item in [
        "Create unsupported RAG scores, forecasts, or ROI calculations",
        "Invent owners, dates, risks, causes, benefits, or approvals",
        "Evaluate individual performance",
        "Silently change source information",
        "Automatically resolve conflicting information",
        "Claim that reminders have been sent",
        "Provide final management, legal, audit, or assurance conclusions",
    ]:
        bullet(doc, item)

    heading(doc, "6. Response Structure")
    for item in [
        "Verified Source Facts",
        "Generated Observations",
        "Decisions Required",
        "Data Gaps",
        "Proposed Actions (labelled: Proposed - requires owner confirmation.)",
    ]:
        bullet(doc, item)

    heading(doc, "7. Acceptance Criteria")
    for item in [
        "Important statements are traceable to approved knowledge.",
        "Source-reported status, dates, amounts, and owners are preserved.",
        "Missing information is identified rather than invented.",
        "Conflicting source values are shown—not silently resolved.",
        "Overdue actions use recorded due dates only.",
        "Financial outputs preserve currency and reporting period.",
        "AI-generated actions are labelled as proposed.",
        "Governance outputs include a professional-review notice.",
    ]:
        bullet(doc, item)

    save_docx(doc, LAB02 / "Program_Portfolio_Intelligence_Agent_BRD.docx")


def create_lab02_program_overview() -> None:
    doc = Document()
    heading(doc, "Horizon Transformation Program — Program Overview")
    para(doc, "Program: Horizon Transformation Program")
    para(doc, "Program Director: Elena Vasquez")
    para(doc, "PMO Lead: Robert Singh")
    para(doc, f"Reporting Period: Q3 2026 (as of {TODAY.strftime('%d %B %Y')})")
    doc.add_paragraph()

    heading(doc, "Program Scope")
    para(
        doc,
        "The Horizon Transformation Program delivers three strategic workstreams: Project Alpha "
        "(Core ERP rollout), Project Beta (Customer portal modernization), and Project Gamma "
        "(Data platform uplift).",
    )

    heading(doc, "Overall Program Status")
    para(doc, "Overall program status (source reported): Amber")
    bullet(doc, "Project Alpha: Amber — integration testing in progress.")
    bullet(doc, "Project Beta: Green — on track for UAT.")
    bullet(doc, "Project Gamma: Amber — dependency on Alpha data migration.")

    heading(doc, "Key Milestones")
    bullet(doc, "Project Alpha UAT complete — Target: 15 October 2026")
    bullet(doc, "Project Beta go-live readiness review — Target: 1 November 2026")
    bullet(doc, "Project Gamma platform cutover — Target: 31 January 2027")

    save_docx(doc, LAB02 / "01_Program_Overview.docx")


def create_lab02_status_updates() -> None:
    doc = Document()
    heading(doc, "Project Status Updates — Week Ending 19 September 2026")

    heading(doc, "Project Alpha — Core ERP Rollout", 2)
    para(doc, "Status: Amber")
    para(doc, "Project Manager: John Smith")
    bullet(doc, "Integration testing 78% complete.")
    bullet(doc, "Data migration rehearsal completed successfully.")
    bullet(doc, "Open issue: vendor patch required for payroll interface.")
    para(doc, "Planned go-live date: 30 November 2026")

    heading(doc, "Project Beta — Customer Portal Modernization", 2)
    para(doc, "Status: Green")
    para(doc, "Project Manager: Aisha Rahman")
    bullet(doc, "UAT started on schedule.")
    bullet(doc, "All critical defects resolved.")
    para(doc, "Planned go-live date: 1 December 2026")

    heading(doc, "Project Gamma — Data Platform Uplift", 2)
    para(doc, "Status: Amber")
    para(doc, "Project Manager: Carlos Mendez")
    bullet(doc, "Platform provisioning complete.")
    bullet(doc, "Waiting on Alpha data migration dependency.")
    para(doc, "Planned go-live date: 31 January 2027")

    save_docx(doc, LAB02 / "02_Project_Status_Updates.docx")


def create_lab02_steering_minutes() -> None:
    doc = Document()
    heading(doc, "Steering Committee Minutes")
    para(doc, "Program: Horizon Transformation Program")
    para(doc, "Meeting Date: 12 September 2026")
    para(doc, "Chair: Elena Vasquez")
    doc.add_paragraph()

    heading(doc, "Attendees")
    bullet(doc, "Elena Vasquez — Program Director")
    bullet(doc, "Robert Singh — PMO Lead")
    bullet(doc, "John Smith — Project Alpha PM")
    bullet(doc, "Aisha Rahman — Project Beta PM")
    bullet(doc, "Carlos Mendez — Project Gamma PM")
    bullet(doc, "Finance Representative: Helen Grant")

    heading(doc, "Decisions Made")
    bullet(doc, "Approved additional testing window for Project Alpha payroll interface.")
    bullet(doc, "Confirmed Project Beta UAT entry criteria met.")

    heading(doc, "Decisions Required")
    bullet(doc, "Approve revised go-live date for Project Alpha.")
    bullet(doc, "Confirm funding reallocation for Gamma acceleration (deferred).")

    heading(doc, "Discussion Notes — Project Alpha Go-Live")
    para(
        doc,
        "Steering Committee reviewed vendor readiness and cutover planning. Program leadership "
        "noted that a revised go-live of 15 December 2026 may be required to accommodate "
        "payroll interface remediation and extended regression testing.",
    )
    para(doc, "Recorded planned go-live date (Steering Committee): 15 December 2026")

    heading(doc, "Actions from Meeting")
    bullet(doc, "PMO to circulate updated cutover plan — Owner: Robert Singh — Due: 26 September 2026")

    save_docx(doc, LAB02 / "07_Steering_Committee_Minutes.docx")


def create_lab02_raid() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "RAID Register"
    headers = ["ID", "Type", "Project", "Description", "Owner", "Status", "Date Raised", "Target Date"]
    rows = [
        ["R-01", "Risk", "Project Alpha", "Vendor payroll patch may delay cutover", "John Smith", "Open", "2026-08-01", "2026-10-15"],
        ["R-02", "Risk", "Project Gamma", "Data migration dependency on Alpha", "Carlos Mendez", "Open", "2026-07-20", "2026-11-01"],
        ["I-01", "Issue", "Project Alpha", "Payroll interface defect in test environment", "John Smith", "Open", "2026-09-05", "2026-09-30"],
        ["I-02", "Issue", "Project Beta", "Legacy SSO certificate expiring", "Aisha Rahman", "Closed", "2026-08-12", "2026-09-01"],
        ["A-01", "Assumption", "Program", "Vendor delivers patch by 30 September 2026", "Robert Singh", "Active", "2026-08-01", ""],
        ["D-01", "Dependency", "Project Gamma", "Requires Alpha data migration completion", "Carlos Mendez", "Open", "2026-06-01", "2026-11-30"],
        ["D-02", "Dependency", "Project Beta", "Requires network firewall rule approval", "Aisha Rahman", "Open", "2026-07-15", "2026-10-01"],
        ["D-03", "Dependency", "Project Alpha", "Requires HR policy sign-off before payroll go-live", "John Smith", "Open", "2026-05-10", "2026-10-20"],
        ["D-04", "Dependency", "Program", "Shared test environment availability", "Robert Singh", "Open", "2026-04-01", "2026-09-30"],
    ]
    write_sheet(ws, headers, rows)
    wb.save(LAB02 / "03_RAID_Register.xlsx")


def create_lab02_actions() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Action Register"
    headers = ["Action ID", "Action", "Owner", "Due Date", "Status", "Project / Workstream"]
    rows = [
        ["ACT-001", "Submit vendor patch impact assessment", "John Smith", "2026-09-05", "Overdue", "Project Alpha"],
        ["ACT-002", "Circulate updated cutover plan", "Robert Singh", "2026-09-26", "Open", "Program"],
        ["ACT-003", "Complete firewall rule approval", "Aisha Rahman", "2026-09-10", "Overdue", "Project Beta"],
        ["ACT-004", "Confirm HR policy sign-off", "John Smith", "2026-09-15", "Overdue", "Project Alpha"],
        ["ACT-005", "Finalize Gamma migration prerequisites", "Carlos Mendez", "2026-10-05", "Open", "Project Gamma"],
        ["ACT-006", "Publish weekly status dashboard", "Robert Singh", "2026-09-22", "Open", "Program"],
    ]
    write_sheet(ws, headers, rows)
    wb.save(LAB02 / "04_Action_Register.xlsx")


def create_lab02_decisions() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Decision Register"
    headers = ["Decision ID", "Decision", "Status", "Project", "Required By", "Decision Maker", "Notes"]
    rows = [
        ["DEC-001", "Approve additional testing window for Alpha payroll interface", "Made", "Project Alpha", "2026-09-12", "Steering Committee", "Recorded in 12 Sep minutes"],
        ["DEC-002", "Approve revised Project Alpha go-live date", "Required", "Project Alpha", "2026-09-26", "Steering Committee", "Conflict between status update and minutes"],
        ["DEC-003", "Confirm Beta UAT entry criteria met", "Made", "Project Beta", "2026-09-12", "Steering Committee", ""],
        ["DEC-004", "Approve Gamma acceleration funding", "Required", "Project Gamma", "2026-10-15", "Steering Committee", "Deferred at last meeting"],
    ]
    write_sheet(ws, headers, rows)
    wb.save(LAB02 / "05_Decision_Register.xlsx")


def create_lab02_financials() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Financial Summary"
    headers = [
        "Project",
        "Reporting Period",
        "Currency",
        "Budget",
        "Actual",
        "Forecast",
        "Variance",
        "Notes",
    ]
    rows = [
        ["Project Alpha", "Q3 2026", "USD", 4200000, 3150000, 4450000, 250000, "Forecast includes vendor patch costs"],
        ["Project Beta", "Q3 2026", "USD", 1800000, 920000, 1750000, -50000, "Under budget due to deferred licensing"],
        ["Project Gamma", "Q3 2026", "USD", 2500000, 680000, 2600000, 100000, "Spend ramp expected in Q4"],
        ["Program Total", "Q3 2026", "USD", 8500000, 4750000, 8800000, 300000, "Program-level rollup"],
    ]
    write_sheet(ws, headers, rows)

    ws2 = wb.create_sheet("Notes")
    write_sheet(
        ws2,
        ["Field", "Availability"],
        [
            ["ROI", "Not provided in approved financial records"],
            ["Benefits realization", "Not provided in approved financial records"],
        ],
    )
    wb.save(LAB02 / "06_Financial_Summary.xlsx")


def create_lab03_agent_brd() -> None:
    doc = Document()
    title = doc.add_heading("Sprint Planning Agent — Business Requirements Document", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    para(doc, "Document Version: 1.0  |  Status: Draft for Training  |  Classification: Fictional / Sanitized")
    doc.add_paragraph()

    heading(doc, "1. Purpose")
    para(
        doc,
        "Define requirements for a Sprint Planning Agent that helps agile teams convert an "
        "approved product BRD into a structured sprint plan draft using delivery policy, "
        "best practices, team capacity, and Definition of Done.",
    )

    heading(doc, "2. Primary Users")
    for user in ["Scrum Master", "Product Owner", "Development Team", "Delivery Manager", "Business Analyst"]:
        bullet(doc, user)

    heading(doc, "3. Approved Knowledge Sources")
    for src in [
        "Product BRD (Customer Self-Service Portal)",
        "Agile Delivery Policy",
        "Scrum Best Practices Guide",
        "Team Capacity Planner",
        "Definition of Done",
    ]:
        bullet(doc, src)

    heading(doc, "4. Functional Requirements")
    reqs = [
        ("SPR-01", "Analyze the product BRD and summarize requirements by Must/Should/Could priority."),
        ("SPR-02", "Propose a sprint goal based on Must and Should requirements."),
        ("SPR-03", "Create user stories with acceptance criteria only when supported by the BRD."),
        ("SPR-04", "Build a proposed sprint backlog with Fibonacci story point estimates."),
        ("SPR-05", "Respect recorded team capacity from the Team Capacity Planner."),
        ("SPR-06", "Apply Agile Delivery Policy limits including maximum backlog items per sprint."),
        ("SPR-07", "Reference Definition of Done when describing story completion criteria."),
        ("SPR-08", "Identify missing, incomplete, or ambiguous acceptance criteria."),
        ("SPR-09", "Show conflicting requirements without silently resolving them."),
        ("SPR-10", "Label proposed estimates and plans as drafts requiring team review."),
    ]
    for req_id, text in reqs:
        p = doc.add_paragraph(style="List Number")
        p.add_run(f"{req_id}: ").bold = True
        p.add_run(text)

    heading(doc, "5. Controls — Agent Must NOT")
    for item in [
        "Invent requirements or acceptance criteria not in the BRD",
        "Invent historical team velocity",
        "Exceed recorded capacity without flagging overrun",
        "Commit more backlog items than policy allows",
        "Silently resolve conflicting requirements",
        "Present the sprint plan as approved or committed",
        "Include Could-have requirements without Product Owner confirmation",
    ]:
        bullet(doc, item)

    heading(doc, "6. Acceptance Criteria")
    for item in [
        "Stories trace to BRD requirement IDs.",
        "Acceptance criteria match BRD text or are marked missing.",
        "Sprint backlog respects 34 story point capacity and 7 item limit.",
        "REQ-05 and REQ-07 gaps are identified.",
        "REQ-08 / REQ-09 notification conflict is shown.",
        "Estimates are labelled as proposed.",
        "Output includes professional-review notice.",
    ]:
        bullet(doc, item)

    save_docx(doc, LAB03 / "Sprint_Planning_Agent_BRD.docx")


def create_lab03_product_brd() -> None:
    doc = Document()
    title = doc.add_heading("Customer Self-Service Portal — Product BRD", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    para(doc, "Product: Customer Self-Service Portal  |  Sprint: Sprint 4  |  Version: 1.0  |  Fictional / Sanitized")
    doc.add_paragraph()

    heading(doc, "1. Overview")
    para(
        doc,
        "This BRD defines features for the Customer Self-Service Portal enabling customers "
        "to manage profiles, view orders, submit support requests, and configure notifications.",
    )

    heading(doc, "2. Requirements")

    requirements = [
        (
            "REQ-01",
            "Must",
            "User Registration",
            "Customers can register using email and verify their account.",
            "1. User can enter email and password.\n2. Verification email is sent within 60 seconds.\n3. Account activates only after email verification.",
        ),
        (
            "REQ-02",
            "Must",
            "Profile Management",
            "Customers can view and update their profile information.",
            "1. User can update name, phone, and address.\n2. Changes save successfully and display confirmation.\n3. Invalid phone format shows validation error.",
        ),
        (
            "REQ-03",
            "Must",
            "Order History",
            "Customers can view a list of past orders with status.",
            "1. Order list shows order ID, date, status, and total.\n2. List supports pagination for more than 20 orders.\n3. Empty state message displays when no orders exist.",
        ),
        (
            "REQ-04",
            "Must",
            "Support Ticket Submission",
            "Customers can submit a support ticket from the portal.",
            "1. User can enter subject, category, and description.\n2. Ticket confirmation displays a ticket reference number.\n3. Ticket appears in the user's ticket history.",
        ),
        (
            "REQ-05",
            "Should",
            "Live Chat Integration",
            "Customers can initiate live chat with a support agent during business hours.",
            "",
        ),
        (
            "REQ-06",
            "Must",
            "Password Reset via SMS",
            "Customers can reset their password using SMS verification.",
            "1. User can request reset via registered mobile number.\n2. SMS code expires after 10 minutes.\n3. Successful reset allows login with new password.",
        ),
        (
            "REQ-07",
            "Should",
            "Multi-Language Support",
            "Portal supports major languages for customer-facing pages.",
            "",
        ),
        (
            "REQ-08",
            "Must",
            "Email Notification Preferences",
            "Customers can configure notification preferences. Default notification channel: Email.",
            "1. User can enable or disable email notifications by category.\n2. Default channel on first login is Email.\n3. Preferences persist after logout.",
        ),
        (
            "REQ-09",
            "Must",
            "Notification Delivery Channel",
            "System sends customer notifications through the configured delivery channel. Default notification delivery channel: SMS.",
            "1. Notifications use the configured default channel.\n2. User can view current default channel in settings.\n3. Channel change requires confirmation.",
        ),
        (
            "REQ-10",
            "Could",
            "Export Order History PDF",
            "Customers can export order history as a PDF document.",
            "1. User can export last 12 months of orders.\n2. PDF includes order ID, date, status, and total.\n3. Export completes within 30 seconds for standard accounts.",
        ),
        (
            "REQ-11",
            "Must",
            "Accessibility WCAG 2.1 AA",
            "Portal meets WCAG 2.1 AA accessibility standards for core customer journeys.",
            "1. All core flows pass automated accessibility scan.\n2. Keyboard navigation works for registration, login, and ticket submission.\n3. Color contrast meets WCAG 2.1 AA minimum.",
        ),
        (
            "REQ-12",
            "Could",
            "Loyalty Points Dashboard",
            "Customers can view loyalty points balance and recent activity.",
            "1. Dashboard shows current points balance.\n2. Recent activity lists last 10 transactions.\n3. Empty state displays when no loyalty activity exists.",
        ),
    ]

    for req_id, priority, title_text, description, ac in requirements:
        heading(doc, f"{req_id} — {title_text} ({priority})", 2)
        para(doc, f"Priority: {priority}")
        para(doc, description)
        heading(doc, "Acceptance Criteria", 3)
        if ac:
            for line in ac.split("\n"):
                bullet(doc, line)
        else:
            para(doc, "Acceptance criteria: Not provided in this BRD.")

    save_docx(doc, LAB03 / "Customer_Self_Service_Portal_BRD.docx")


def create_lab03_agile_policy() -> None:
    doc = Document()
    heading(doc, "Agile Delivery Policy — Contoso Digital")
    para(doc, "Policy Version: 2.1  |  Effective: January 2026  |  Applies to: All software delivery teams")
    doc.add_paragraph()

    heading(doc, "1. Sprint Structure")
    bullet(doc, "Sprint length: 2 weeks (10 working days).")
    bullet(doc, "Sprint numbering follows product team calendar.")
    bullet(doc, "Sprint 4 is the current planning sprint for Customer Self-Service Portal.")

    heading(doc, "2. Sprint Planning Rules")
    bullet(doc, "Every sprint must have one sprint goal approved by the Product Owner.")
    bullet(doc, "Maximum 7 product backlog items may be committed per sprint.")
    bullet(doc, "Could-have requirements require explicit Product Owner approval to enter a sprint.")
    bullet(doc, "All committed stories must map to an approved requirement ID.")

    heading(doc, "3. Estimation")
    bullet(doc, "Use Fibonacci story points: 1, 2, 3, 5, 8, 13, 21.")
    bullet(doc, "Story points represent relative effort, not hours.")
    bullet(doc, "Estimates require team consensus during sprint planning.")

    heading(doc, "4. Capacity")
    bullet(doc, "Sprint capacity must be taken from the Team Capacity Planner for the sprint.")
    bullet(doc, "Historical velocity may be used only when recorded in approved team artifacts.")
    bullet(doc, "Plans exceeding recorded capacity require explicit team discussion.")

    heading(doc, "5. Sprint Ceremonies")
    bullet(doc, "Sprint Planning — Day 1 (max 4 hours)")
    bullet(doc, "Daily Stand-up — 15 minutes each working day")
    bullet(doc, "Sprint Review — Last day (2 hours)")
    bullet(doc, "Sprint Retrospective — Last day (1.5 hours)")

    heading(doc, "6. Definition of Done")
    para(doc, "All stories must meet the team Definition of Done before being accepted in a sprint review.")

    save_docx(doc, LAB03 / "Agile_Delivery_Policy.docx")


def create_lab03_scrum_guide() -> None:
    doc = Document()
    heading(doc, "Scrum Best Practices Guide — Contoso Digital")
    doc.add_paragraph()

    heading(doc, "1. User Story Format")
    para(doc, 'Use: "As a [user type], I want [goal], so that [benefit]."')

    heading(doc, "2. INVEST Criteria")
    for item in ["Independent", "Negotiable", "Valuable", "Estimable", "Small", "Testable"]:
        bullet(doc, item)

    heading(doc, "3. Story Splitting")
    bullet(doc, "Split by workflow step when a requirement covers multiple user actions.")
    bullet(doc, "Split by business rule variation when criteria differ by scenario.")
    bullet(doc, "Label proposed splits as requiring Product Owner confirmation.")

    heading(doc, "4. Acceptance Criteria")
    bullet(doc, "Copy acceptance criteria from the BRD when available.")
    bullet(doc, "Do not invent acceptance criteria—request Product Owner clarification instead.")
    bullet(doc, "Each acceptance criterion should be testable.")

    heading(doc, "5. Estimation Guidance")
    bullet(doc, "Compare new stories to previously completed stories of similar complexity.")
    bullet(doc, "If no historical reference exists in approved sources, label estimate as proposed.")
    bullet(doc, "Stories larger than 13 points should be split before sprint commitment.")

    heading(doc, "6. Sprint Goal")
    bullet(doc, "Sprint goal should reflect Must and Should requirements for the sprint.")
    bullet(doc, "Goal should be understandable to business stakeholders in one sentence.")

    save_docx(doc, LAB03 / "Scrum_Best_Practices_Guide.docx")


def create_lab03_capacity() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sprint 4 Summary"
    write_sheet(
        ws,
        ["Field", "Value", "Notes"],
        [
            ["Product", "Customer Self-Service Portal", ""],
            ["Sprint", "Sprint 4", "2-week sprint"],
            ["Sprint Capacity (Story Points)", 34, "Recorded planning capacity"],
            ["Maximum Backlog Items", 7, "Per Agile Delivery Policy"],
            ["Historical Velocity", "Not provided", "Not available in approved artifacts"],
        ],
    )

    ws2 = wb.create_sheet("Team Roster")
    write_sheet(
        ws2,
        ["Team Member", "Role", "Availability (%)", "Notes"],
        [
            ["Alex Morgan", "Developer", 100, ""],
            ["Jordan Lee", "Developer", 100, ""],
            ["Sam Patel", "Developer", 80, "Leave days 3–4"],
            ["Taylor Kim", "Developer", 100, ""],
            ["Casey Nguyen", "Developer", 100, ""],
            ["Riley Brooks", "QA Engineer", 100, ""],
            ["Morgan Hayes", "Product Owner", 50, "Shared across two products"],
            ["Jamie Ortiz", "Scrum Master", 30, "Facilitation only"],
        ],
    )
    wb.save(LAB03 / "Team_Capacity_Planner.xlsx")


def create_lab03_dod() -> None:
    doc = Document()
    heading(doc, "Definition of Done — Customer Self-Service Portal Team")
    para(doc, "Team: CSS Portal Squad  |  Version: 1.3  |  Applies to: All sprint stories")
    doc.add_paragraph()

    heading(doc, "Story-Level Definition of Done")
    items = [
        "Code complete and peer reviewed.",
        "Unit tests written and passing.",
        "Integration tests passing in dev environment.",
        "Acceptance criteria verified by QA.",
        "Product Owner acceptance obtained.",
        "No critical or high defects open.",
        "Documentation updated if API or user flow changed.",
        "Accessibility check completed for customer-facing UI changes.",
        "Deployed to test environment.",
        "Demo-ready for Sprint Review.",
    ]
    for item in items:
        bullet(doc, item)

    heading(doc, "Notes")
    bullet(doc, "Stories without BRD acceptance criteria cannot be marked Done until PO provides them.")
    bullet(doc, "Accessibility verification required for all Must and Should requirements affecting UI.")

    save_docx(doc, LAB03 / "Definition_of_Done.docx")


def main() -> None:
    print("Generating Lab 01 reference files...")
    create_lab01_brd()
    create_lab01_data()

    print("Generating Lab 02 reference files...")
    create_lab02_brd()
    create_lab02_program_overview()
    create_lab02_status_updates()
    create_lab02_steering_minutes()
    create_lab02_raid()
    create_lab02_actions()
    create_lab02_decisions()
    create_lab02_financials()

    print("Generating Lab 03 reference files...")
    create_lab03_agent_brd()
    create_lab03_product_brd()
    create_lab03_agile_policy()
    create_lab03_scrum_guide()
    create_lab03_capacity()
    create_lab03_dod()

    print(f"Done. Files written to: {ROOT}")


if __name__ == "__main__":
    main()
